"""L1 单元格归一化（TECH_DESIGN §5.1、§7.2）。

把 xlsx 读成与 Excel 行列一一对应的网格：值、单元格标记（隐藏/删除线/错误值/公式无缓存值/合并填充/加粗）。
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import warnings
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from typing import Any

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

# 单元格标记位
HIDDEN_ROW = 1 << 0
HIDDEN_COL = 1 << 1
STRIKE = 1 << 2
ERROR = 1 << 3
F_NOCACHE = 1 << 4
MERGED_FILLED = 1 << 5
BOLD = 1 << 6

EXCEL_ERRORS = {"#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!", "#GETTING_DATA", "#SPILL!", "#CALC!"}


class WorkbookError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def to_a1(row: int, col: int) -> str:
    """1-based 行列 → A1 地址。"""
    return f"{get_column_letter(col)}{row}"


_A1 = re.compile(r"^([A-Z]+)(\d+)$")


def from_a1(addr: str) -> tuple[int, int]:
    m = _A1.match(addr.upper())
    if not m:
        raise ValueError(addr)
    letters, row = m.groups()
    col = 0
    for ch in letters:
        col = col * 26 + (ord(ch) - 64)
    return int(row), col


_WS = re.compile(r"\s+")
_INVISIBLE = dict.fromkeys(map(ord, "​‌‍﻿⁠"), None)


def norm_text(v: Any) -> str:
    """表头/锚点文本归一化：全半角统一、去不可见字符、合并空白、去首尾空白。"""
    if v is None:
        return ""
    s = str(v).translate(_INVISIBLE)
    s = unicodedata.normalize("NFKC", s)
    return _WS.sub(" ", s).strip()


def is_blank(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str):
        return norm_text(v) == ""
    if isinstance(v, float) and np.isnan(v):
        return True
    return False


@dataclass
class SheetGrid:
    name: str
    values: np.ndarray  # object，0-based；values[r-1, c-1] 对应 Excel (r, c)
    flags: np.ndarray  # uint16 标记位
    number_formats: np.ndarray
    date1904: bool
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # (r1, c1, r2, c2) 1-based
    outline: dict[int, int] = field(default_factory=dict)  # 行号 → 大纲级别

    @property
    def n_rows(self) -> int:
        return self.values.shape[0]

    @property
    def n_cols(self) -> int:
        return self.values.shape[1]

    def cell(self, row: int, col: int) -> Any:
        if 1 <= row <= self.n_rows and 1 <= col <= self.n_cols:
            return self.values[row - 1, col - 1]
        return None

    def flag(self, row: int, col: int) -> int:
        if 1 <= row <= self.n_rows and 1 <= col <= self.n_cols:
            return int(self.flags[row - 1, col - 1])
        return 0

    def fingerprint(self) -> str:
        """结构指纹：前 10 个非空行的归一化文本。"""
        h = hashlib.sha256(self.name.encode())
        seen = 0
        for r in range(self.n_rows):
            row = [norm_text(v) for v in self.values[r] if not is_blank(v)]
            if row:
                h.update("|".join(row).encode())
                seen += 1
                if seen >= 10:
                    break
        return h.hexdigest()[:16]


def _rich_to_plain(v: Any) -> Any:
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    # CellRichText、ArrayFormula 等对象
    text = getattr(v, "text", None)
    if isinstance(text, str):
        return text
    if hasattr(v, "__iter__") and not hasattr(v, "year"):
        try:
            return "".join(str(getattr(p, "text", p)) for p in v)
        except TypeError:
            return str(v)
    return v


_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@dataclass
class _SheetXmlInfo:
    merged: list[tuple[int, int, int, int]]
    hidden_rows: set[int]
    hidden_cols: set[int]
    outline: dict[int, int]
    formula_no_cache: set[tuple[int, int]]


def _scan_sheet_xml(raw: bytes) -> _SheetXmlInfo:
    """一次流式扫描工作表 XML：合并单元格、隐藏行列、大纲级别、「有公式但无缓存值」的单元格。"""
    info = _SheetXmlInfo([], set(), set(), {}, set())
    cell_ref, has_f, has_v = None, False, False
    for event, el in ET.iterparse(io.BytesIO(raw), events=("start", "end")):
        tag = el.tag[len(_NS):] if el.tag.startswith(_NS) else el.tag
        if event == "start":
            if tag == "c":
                cell_ref, has_f, has_v = el.get("r"), False, False
            elif tag == "row":
                r = el.get("r")
                if r is not None:
                    if el.get("hidden") in ("1", "true"):
                        info.hidden_rows.add(int(r))
                    lvl = el.get("outlineLevel")
                    if lvl and lvl != "0":
                        info.outline[int(r)] = int(lvl)
            elif tag == "col":
                if el.get("hidden") in ("1", "true"):
                    for c in range(int(el.get("min", "0")), int(el.get("max", "0")) + 1):
                        info.hidden_cols.add(c)
            elif tag == "mergeCell":
                ref = el.get("ref", "")
                if ":" in ref:
                    a, b = ref.split(":")
                    r1, c1 = from_a1(a)
                    r2, c2 = from_a1(b)
                    info.merged.append((r1, c1, r2, c2))
        else:
            if tag == "f":
                has_f = True
            elif tag == "v" and (el.text or "") != "":
                has_v = True
            elif tag == "c":
                if has_f and not has_v and cell_ref:
                    info.formula_no_cache.add(from_a1(cell_ref))
                el.clear()
            elif tag == "row":
                el.clear()
    return info


MAX_UNCOMPRESSED_BYTES = 512 * 2**20


class Workbook:
    """一个上传文件：按需加载 Sheet 网格（只读流式），同一 Sheet 只加载一次。"""

    def __init__(self, data: bytes, max_cells: int | None = None):
        self._data = data
        self.max_cells = max_cells
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            try:
                self._zip = zipfile.ZipFile(io.BytesIO(data))
                self._check_zip_bomb()
                self._wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True, keep_vba=False)
            except WorkbookError:
                raise
            except Exception as exc:
                raise WorkbookError("FILE_UNSUPPORTED", f"无法读取文件，请确认是 xlsx 格式：{type(exc).__name__}") from exc
        self.date1904 = bool(getattr(self._wb, "epoch", None) and self._wb.epoch.year == 1904)
        self._states = self._sheet_states()
        self._cache: dict[tuple[str, str], SheetGrid] = {}

    def _check_zip_bomb(self) -> None:
        """解压炸弹防护：按 zip 目录声明的解压后大小判断（zipfile 读取时不会超过声明的大小）。
        正常 xlsx 压缩比约 5~10 倍，上限按 20MB 文件 × 25 倍留足余量。"""
        total = sum(i.file_size for i in self._zip.infolist())
        if total > MAX_UNCOMPRESSED_BYTES:
            raise WorkbookError("FILE_TOO_LARGE", f"文件解压后超过 {MAX_UNCOMPRESSED_BYTES // 2**20}MB，疑似异常文件")

    def _sheet_states(self) -> dict[str, str]:
        try:
            root = ET.fromstring(self._zip.read("xl/workbook.xml"))
            return {s.get("name"): s.get("state", "visible") for s in root.iter(_NS + "sheet")}
        except Exception:
            return {}

    @property
    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def sheet(self, name: str, merge_policy: str = "FILL") -> SheetGrid:
        key = (name, merge_policy)
        if key not in self._cache:
            if name not in self._wb.sheetnames:
                raise WorkbookError("SHEET_NOT_FOUND", f"找不到 Sheet「{name}」，现有：{', '.join(self.sheet_names)}")
            self._cache[key] = self._load(name, merge_policy)
        return self._cache[key]

    def _load(self, name: str, merge_policy: str) -> SheetGrid:
        ws = self._wb[name]
        path = getattr(ws, "_worksheet_path", None)
        info = _scan_sheet_xml(self._zip.read(path)) if path else _SheetXmlInfo([], set(), set(), {}, set())
        rows: list[list] = []
        max_c, max_r = 0, 0
        for row in ws.iter_rows():
            vals = []
            for c in row:
                if getattr(c, "value", None) is None and not hasattr(c, "row"):
                    vals.append((None, 0, None))
                    continue
                v = _rich_to_plain(c.value)
                f = 0
                if v is not None and getattr(c, "has_style", False):
                    font = c.font
                    if font is not None:
                        if font.strike:
                            f |= STRIKE
                        if font.b:
                            f |= BOLD
                vals.append((v, f, c.number_format if v is not None else None))
                if v is not None and not (isinstance(v, str) and v.strip() == ""):
                    max_c = max(max_c, len(vals))
                    max_r = len(rows) + 1
            rows.append(vals)
        for (r1, c1, _r2, _c2) in info.merged:
            max_r, max_c = max(max_r, r1), max(max_c, c1)
        for (r, c) in info.formula_no_cache:
            max_r, max_c = max(max_r, r), max(max_c, c)
        h, w = max(max_r, 1), max(max_c, 1)
        if self.max_cells and h * w > self.max_cells:
            raise WorkbookError("FILE_TOO_LARGE", f"Sheet「{name}」有 {h}×{w} 个单元格，超过上限 {self.max_cells}")
        values = np.empty((h, w), dtype=object)
        flags = np.zeros((h, w), dtype=np.uint16)
        fmts = np.empty((h, w), dtype=object)
        for ri in range(min(h, len(rows))):
            row = rows[ri]
            for ci in range(min(w, len(row))):
                v, f, fmt = row[ci]
                if isinstance(v, str) and v.strip() in EXCEL_ERRORS:
                    f |= ERROR
                values[ri, ci], flags[ri, ci], fmts[ri, ci] = v, f, fmt
        for (r, c) in info.formula_no_cache:
            if values[r - 1, c - 1] is None:
                flags[r - 1, c - 1] |= F_NOCACHE
        for r in info.hidden_rows:
            if r <= h:
                flags[r - 1, :] |= HIDDEN_ROW
        for c in info.hidden_cols:
            if c <= w:
                flags[:, c - 1] |= HIDDEN_COL
        merged = []
        for (r1, c1, r2, c2) in info.merged:
            r2, c2 = min(r2, h), min(c2, w)
            merged.append((r1, c1, r2, c2))
            if merge_policy == "FILL":
                top_left = values[r1 - 1, c1 - 1]
                values[r1 - 1 : r2, c1 - 1 : c2] = top_left
                flags[r1 - 1 : r2, c1 - 1 : c2] |= MERGED_FILLED
                flags[r1 - 1, c1 - 1] &= ~np.uint16(MERGED_FILLED)
        outline = {r: lvl for r, lvl in info.outline.items() if r <= h}
        return SheetGrid(name, values, flags, fmts, self.date1904, merged, outline)

    def sheet_meta(self, names: list[str] | None = None) -> list[dict]:
        out = []
        for n in names or self.sheet_names:
            g = self.sheet(n)
            out.append(
                {
                    "name": n,
                    "maxRow": g.n_rows,
                    "maxCol": g.n_cols,
                    "fingerprint": g.fingerprint(),
                    "hasUncachedFormula": bool((g.flags & F_NOCACHE).any()),
                    "hidden": self._states.get(n, "visible") != "visible",
                }
            )
        return out
